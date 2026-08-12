"""Write the .xlsx sources for the formula test fixtures.

Excel has to do the .xlsb conversion itself -- see files/README.md. Usage:

    python tests/generate_fixtures.py <output-dir>
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

OUTDIR = sys.argv[1] if len(sys.argv) > 1 else '.'


def build_corpus():
  FORMULAS = [
      # --- operators ---
      '=A1+A2', '=A1-A2', '=A1*A2', '=A1/A2', '=A1^A2', '=A1&A2',
      '=A1<A2', '=A1<=A2', '=A1=A2', '=A1>=A2', '=A1>A2', '=A1<>A2',
      '=-A1', '=+A1', '=A1%', '=(A1+A2)*A3', '=A1+A2*A3', '=(A1+A2)/(A3-A4)',
      # --- constants ---
      '=1', '=1.5', '=-1.5', '=1000000', '="hello"', '="quote""inside"',
      '=TRUE', '=FALSE', '=#REF!', '=#DIV/0!', '=#N/A',
      # --- reference forms / absolute flags ---
      '=A1', '=$A$1', '=$A1', '=A$1', '=A1:A5', '=$A$1:$A$5', '=A$1:$A5',
      '=SUM(A1:A5)', '=SUM($A$1:$A$5)', '=SUM(A:A)', '=SUM(1:1)',
      # --- ranges / set operators ---
      '=SUM(A1:A3 A2:A5)', '=SUM(A1:A2,A4:A5)', '=SUM((A1:A2,A4:A5))',
      '=SUM(A1:A5,A1)', '=COUNT(A1:A5)',
      # --- functions: fixed arity, variable arity, nested ---
      '=MAX(A1:A5)', '=MIN(A1:A5)', '=ROUND(A1/A2,2)', '=ABS(-A1)',
      '=IF(A1>1,"yes","no")', '=IF(A1>1,SUM(A1:A2),MAX(A1:A2))',
      '=IFERROR(A1/0,0)', '=IFERROR(-MAX(SUM(A1:A5),0)*A2,"")',
      '=SUMPRODUCT(A1:A5,A1:A5)', '=COUNTIF(A1:A5,">2")',
      '=SUMIF(A1:A5,">2",A1:A5)', '=VLOOKUP(1,A1:B5,2,FALSE)',
      '=INDEX(A1:A5,2)', '=INDEX(A1:B5,2,2)', '=MATCH(3,A1:A5,0)',
      '=OFFSET(A1,1,1)', '=OFFSET(A1,1,1,2,2)', '=INDIRECT("A1")',
      '=TEXT(A1,"0.00")', '=NOW()', '=TODAY()', '=NA()', '=PI()',
      '=NPV(0.1,A1:A5)', '=IRR(A1:A5)', '=PMT(0.1,10,100)',
      '=CONCATENATE(A1,"x",A2)', '=LEFT("abc",2)', '=ROUND(SUM(A1:A5)/COUNT(A1:A5),4)',
      # --- array constants ---
      '=SUM({1,2,3})', '=SUM({1,2;3,4})', '=SUM({1,"a";TRUE,FALSE})',
      # --- cross-sheet (3D) references ---
      '=Data!A1', '=Data!$A$1', '=SUM(Data!A1:A5)', '=SUM(Data!A1:A5)+A1',
      "='Sheet Space'!A1", "=SUM('Sheet Space'!A1:A5)",
      '=SUM(Data:Data!A1)',
      # --- defined names ---
      '=MyName', '=SUM(MyName)', '=MyName+A1', '=MyCell',
      # --- misc / edge ---
      '=SUM(A1:A5)*2%', '=IF(TRUE,1,2)', '=N(A1)', '=T(A1)',
  ]

  wb = Workbook()

  data = wb.active
  data.title = 'Data'
  for i in range(1, 21):
      data.cell(row=i, column=1, value=float(i))
      data.cell(row=i, column=2, value=float(i * 10))
  data['D1'] = 'alpha'
  data['D2'] = 'beta'

  space = wb.create_sheet('Sheet Space')
  for i in range(1, 11):
      space.cell(row=i, column=1, value=float(i * 100))

  fs = wb.create_sheet('Formulas')
  for i in range(1, 21):
      fs.cell(row=i, column=1, value=float(i))
      fs.cell(row=i, column=2, value=float(i * 3))

  # Formulas go in column D so refs to A/B stay meaningful; E holds the source text.
  for idx, f in enumerate(FORMULAS):
      r = idx + 1
      fs.cell(row=r, column=4, value=f)
      fs.cell(row=r, column=5, value=f.lstrip('='))

  # Shared formulas: Excel collapses a filled-down block into PtgExp hosts.
  shared_start = len(FORMULAS) + 3
  for r in range(shared_start, shared_start + 12):
      fs.cell(row=r, column=4, value='=A{r}*2+B{r}'.format(r=r - shared_start + 1))

  # A table, for structured (PtgList) references.
  tbl_top = shared_start + 15
  fs.cell(row=tbl_top, column=8, value='Qty')
  fs.cell(row=tbl_top, column=9, value='Price')
  for k in range(1, 6):
      fs.cell(row=tbl_top + k, column=8, value=float(k))
      fs.cell(row=tbl_top + k, column=9, value=float(k * 2))
  ref = 'H{}:I{}'.format(tbl_top, tbl_top + 5)
  t = Table(displayName='MyTable', ref=ref)
  t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium9', showRowStripes=True)
  fs.add_table(t)

  struct = [
      '=SUM(MyTable[Qty])',
      '=SUM(MyTable[Price])',
      '=SUMPRODUCT(MyTable[Qty],MyTable[Price])',
  ]
  for k, f in enumerate(struct):
      r = tbl_top + 8 + k
      fs.cell(row=r, column=4, value=f)

  wb.defined_names['MyName'] = DefinedName('MyName', attr_text='Data!$A$1:$A$5')
  wb.defined_names['MyCell'] = DefinedName('MyCell', attr_text='Data!$B$2')
  return wb


def build_extras():
  wb = Workbook()

  data = wb.active
  data.title = "It's Data"
  for i in range(1, 11):
      data.cell(row=i, column=1, value=float(i))
      data.cell(row=i, column=2, value=float(i * 2))

  fs = wb.create_sheet('Calc')
  for i in range(1, 11):
      fs.cell(row=i, column=1, value=float(i))
      fs.cell(row=i, column=2, value=float(i * 3))


  # Single-cell CSE array formulas.
  singles = [
      '=SUM(A1:A5*B1:B5)',
      '=MAX(IF(A1:A5>2,A1:A5))',
      '=SUM(IF(A1:A5>2,1,0))',
  ]
  for k, f in enumerate(singles):
      r = k + 1
      ref = 'D{}'.format(r)
      fs[ref] = ArrayFormula(ref, f)

  # A multi-cell CSE array formula spanning a block.
  fs['F1'] = ArrayFormula('F1:F3', '=TRANSPOSE(A1:C1)')

  # Ordinary formulas that lean on the awkward sheet name and deeper nesting.
  plain = [
      "='It''s Data'!A1",
      "=SUM('It''s Data'!A1:A5)",
      '=IF(A1>1,IF(A2>2,IF(A3>3,"deep","c"),"b"),"a")',
      '=IFERROR(IFERROR(A1/0,A2/0),"both")',
      '=SUM(A1:A5)/IF(COUNT(A1:A5)=0,1,COUNT(A1:A5))',
      '=ROUND(SUMPRODUCT((A1:A5>2)*(B1:B5)),2)',
      '=-(A1+A2)',
      '=--A1',
      '=A1&""&A2',
      '=SUM(A1:A5)-SUM(B1:B5)',
      '=CHOOSE(2,A1,A2,A3)',
      '=SUMIFS(A1:A5,B1:B5,">3",A1:A5,"<5")',
      '=COUNTIFS(A1:A5,">1")',
      '=AVERAGEIF(A1:A5,">2")',
      '=EOMONTH(TODAY(),1)',
      '=YEARFRAC(TODAY(),TODAY()+365)',
      '=XNPV(0.1,A1:A5,A1:A5)',
      '=SUBTOTAL(9,A1:A5)',
      '=AGGREGATE(9,6,A1:A5)',
      '=MyRange',
      '=SUM(MyRange)',
  ]
  for k, f in enumerate(plain):
      r = k + 1
      ref = 'H{}'.format(r)
      fs[ref] = f

  wb.defined_names['MyRange'] = DefinedName('MyRange', attr_text="'It''s Data'!$A$1:$A$5")
  return wb


def build_spaced():
  """Formulas carrying the whitespace their author typed."""
  wb = Workbook()
  ws = wb.active
  ws.title = 'Space'
  for i in range(1, 11):
    ws.cell(row=i, column=1, value=float(i))
    ws.cell(row=i, column=2, value=float(i * 2))
  formulas = [
      '=IFERROR(A1, 0)', '=IFERROR(A1,  0)', '=SUM(A1:A5, A1)',
      '=IF(A1 > 1, "yes", "no")', '=A1 + A2', '=A1  +  A2', '= A1+A2',
      '=A1 %', '=- A1', '=(A1 + A2) * A3', '=SUM( A1:A5 )', '=SUM(A1:A5 )',
      '=ROUND( SUM(A1:A5) / COUNT(A1:A5), 4 )', '=CONCATENATE(A1, " x ", A2)',
      '=IF(A1>1, SUM(A1:A2), MAX(A1:A2))', '="keep  inner  spaces"',
      '=A1&"  "&A2', '=SUM(A1:A3 A2:A5)', '=SUM((A1:A2, A4:A5))',
      '=MAX( A1 , A2 )',
  ]
  for k, f in enumerate(formulas):
    ws.cell(row=k + 1, column=4, value=f)
  return wb


def build_scoped():
  """Sheet-scoped names, multi-sheet spans, and small numeric literals."""
  wb = Workbook()
  alpha = wb.active
  alpha.title = 'Alpha'
  for name in ('Alpha', 'Jan', 'Feb', 'Mar', 'Q1 Data', 'Q4 Data', 'Calc'):
    ws = alpha if name == 'Alpha' else wb.create_sheet(name)
    for i in range(1, 6):
      ws.cell(row=i, column=1, value=float(i))
      ws.cell(row=i, column=2, value=float(i * 2))

  alpha.defined_names.add(DefinedName('LocalRate', attr_text='Alpha!$B$2'))
  wb.defined_names['GlobalRate'] = DefinedName('GlobalRate', attr_text='Alpha!$B$3')

  calc = wb['Calc']
  formulas = [
      '=Alpha!LocalRate', '=Alpha!LocalRate*2', '=GlobalRate',
      '=GlobalRate+Alpha!LocalRate', '=SUM(Jan:Mar!A1)', '=SUM(Jan:Mar!A1:A5)',
      "=SUM('Q1 Data:Q4 Data'!A1)", "='Q1 Data:Q4 Data'!B2",
      '=A1*0.00001', '=A1*0.0000001', '=0.000000123', '=A1/3',
      '=1234567890123', '=A1*1000000',
  ]
  for k, f in enumerate(formulas):
    calc.cell(row=k + 1, column=4, value=f)
  for i in range(1, 6):
    calc.cell(row=i, column=1, value=float(i))
  return wb


def build_deleted():
  """A workbook whose referenced sheet is deleted after the fact.

  `Doomed` has to be removed in Excel once the file is open -- see
  files/README.md -- because the point is what Excel rewrites the references
  into, which is `#REF!`.
  """
  wb = Workbook()
  keep = wb.active
  keep.title = 'Keep'
  doomed = wb.create_sheet('Doomed')
  for i in range(1, 6):
    keep.cell(row=i, column=1, value=float(i))
    doomed.cell(row=i, column=1, value=float(i * 10))
  keep['D1'] = '=Doomed!A1'
  keep['D2'] = '=SUM(Doomed!A1:A5)'
  keep['D3'] = '=IF(A1=1,Doomed!A2,0)'
  keep['D4'] = '=Doomed!A1+A1'
  keep['D5'] = '=A1+A2'
  return wb


if __name__ == '__main__':
  if not os.path.isdir(OUTDIR):
    os.makedirs(OUTDIR)
  for name, build in (('fixture.xlsx', build_corpus),
                      ('fixture2.xlsx', build_extras),
                      ('fixture3.xlsx', build_spaced),
                      ('fixture4.xlsx', build_scoped),
                      ('fixture5.xlsx', build_deleted)):
    path = os.path.join(OUTDIR, name)
    build().save(path)
    print('wrote', path)
